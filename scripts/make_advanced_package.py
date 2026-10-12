"""Build the advanced appendix, offline figure gallery and formatted workbook."""
from __future__ import annotations

import argparse
import html
import io
import json
import math
from pathlib import Path

import fitz
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph

ROOT = Path(__file__).resolve().parents[1]
NAVY = "17354A"
TEAL = "117C83"
INK = "233E50"
GRAY = "526877"
PALE = "EDF5F5"
SHEET_NOTES = {
    "ablation_metrics": "Validation only; five fixed feature sets; 95% patient bootstrap intervals.",
    "ablation_paired_differences": "Paired full minus ablation; positive AP/AUROC and negative Brier favor full.",
    "calibration_bins": "Test reliability bins whose boundaries were fixed on validation.",
    "calibration_comparison": "Paired raw minus calibrated test errors for the frozen selected model.",
    "calibration_summary": "Overall and drug-level test probability checks with patient bootstrap intervals.",
    "drug_performance": "Frozen selected model; test performance at the original global cutoff.",
    "missingness_by_split": "One record per culture, or per culture-candidate for prior candidate resistance.",
    "organism_drug_slices": "Retained label coverage among included organism-events; absent outcomes remain unknown.",
    "probability_histogram": "Distribution of frozen raw/calibrated resistance probabilities.",
    "threshold_operating_points": "Validation-chosen cutoffs applied unchanged to test; counts per 1,000 candidate rows.",
    "validation_tradeoff_curve": "Validation threshold curve; no test-based threshold optimization.",
}


def load_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def plain(value):
    text = str(value)
    if "\u2014" in text:
        raise ValueError("The appendix and gallery must not contain em dashes.")
    return text


def build_content(root, index):
    folder = root / "results" / "advanced"
    ablation = pd.read_csv(folder / "ablation_metrics.csv").set_index(["model", "metric"])
    paired = pd.read_csv(folder / "ablation_paired_differences.csv").set_index(["model_b", "metric"])
    drugs = pd.read_csv(folder / "drug_performance.csv").set_index("candidate_antibiotic")
    comparison = pd.read_csv(folder / "calibration_comparison.csv").set_index("metric")
    calibration = pd.read_csv(folder / "calibration_summary.csv")
    operating = pd.read_csv(folder / "threshold_operating_points.csv")
    operating = operating[operating["split"] == "test"].set_index("operating_point")
    missing = pd.read_csv(folder / "missingness_by_split.csv")
    missing = missing[missing["split"] == "test"].set_index("feature")
    cover = pd.read_csv(folder / "organism_drug_slices.csv").set_index(["organism", "candidate_antibiotic"])
    metadata = load_json(folder / "diagnostics_metadata.json")
    ablation_metadata = load_json(folder / "ablation_metadata.json")
    primary = pd.read_csv(root / "results" / "models" / "model_comparison.csv")
    primary = primary[primary["split"] == "test"].set_index("model")
    ap = lambda model: float(ablation.loc[(model, "antibiotic_macro_auprc_r"), "estimate"])
    delta = paired.loc[("core_prior_history", "antibiotic_macro_auprc_r")]
    value_delta = paired.loc[("core_availability_only", "antibiotic_macro_auprc_r")]
    brier = comparison.loc["brier"]
    original, high = operating.loc["original"], operating.loc["recall_90"]
    pip, amp = drugs.loc["piperacillin/tazobactam"], drugs.loc["ampicillin"]
    pooled = calibration[(calibration.scope == "overall") & (calibration.probability == "calibrated")].iloc[0]
    notes = {
        "01_model_comparison": [
            f"The selected weighted LightGBM has average resistance AP {primary.loc['lightgbm_balanced', 'antibiotic_macro_auprc_r']:.3f} across drugs. The unweighted version scores {primary.loc['lightgbm_unweighted', 'antibiotic_macro_auprc_r']:.3f}. The difference is small; these results do not show a clear benefit from class weighting.",
            "AUROC checks ranking. Resistance average precision (AP) checks how well resistant results rise to the top. These two scores give each drug equal weight. Brier averages probability error over candidate rows. All models and the selected model are unchanged from the original analysis.",
        ],
        "02_feature_group_ablation": [
            f"Core plus prior history scores {ap('core_prior_history'):.3f} for resistance AP; all features score {ap('full'):.3f}. The paired full-minus-history difference is {delta.difference_a_minus_b:.3f} (95% interval {delta.ci_low:.3f} to {delta.ci_high:.3f}). This fixed recipe shows no clear added improvement from labs and vital signs once prior history is included. Similar scores do not prove equivalence.",
            f"Core plus availability patterns scores {ap('core_availability_only'):.3f}. The full model scores {value_delta.difference_a_minus_b:.3f} AP higher (interval {value_delta.ci_low:.3f} to {value_delta.ci_high:.3f}). The control keeps missing values missing and replaces observed numbers with a constant. A recorded prior-resistance entry already conveys history even when its time lag is erased, so this signal is not only about documentation habits.",
        ],
        "03_drug_performance": [
            f"The same global cutoff finds none of the {int(pip.n_r):,} resistant piperacillin/tazobactam rows. It flags just {int(pip.n_flagged):,} row, which is susceptible. For ampicillin it finds every resistant row, but flags all {int(amp.n):,} rows. Overall recall of {original.recall_r:.1%} hides these large differences.",
            "Ranking and cutoff performance answer different questions. Drug-level AP also changes with resistance prevalence and case mix. Wide intervals, such as those for cefepime, show that many rows can still represent relatively few patients with resistant results. A bootstrap interval at zero or one is not certainty about future patients.",
        ],
        "04_overall_calibration": [
            f"Brier falls from {brier.raw:.3f} before calibration to {brier.calibrated:.3f} after calibration. The paired decrease is {brier.raw_minus_calibrated:.3f} (95% interval {brier.raw_minus_calibrated_ci_low:.3f} to {brier.raw_minus_calibrated_ci_high:.3f}). The calibrator was fitted on held-out training patients in the original run, not on these test outcomes.",
            f"The mean calibrated resistance probability is {pooled.mean_predicted_r:.1%}; the observed resistant share is {pooled.prevalence_r:.1%}. The reliability curve checks agreement across probability ranges. The lower Brier score is useful, but by itself does not prove good calibration at every probability.",
        ],
        "05_calibration_by_drug": [
            "The drug panels show where probability checks have enough support and where they do not. Bin boundaries came from validation predictions. Open circles mark fewer than 100 test rows or fewer than 20 resistant results in a bin. Large uncertainty in a sparse bin should remain visible.",
            "An acceptable pooled curve can conceal drug-level errors. Read each panel with its counts and intervals. These are descriptive checks on the existing test set. No drug-specific recalibration was fitted, and the plots do not show that the model will be calibrated at another hospital.",
        ],
        "06_threshold_tradeoffs": [
            f"The original cutoff gives about {original.false_alerts_per_1000:.0f} false resistance alerts and {original.missed_r_per_1000:.0f} missed resistant results per 1,000 test rows. The 90% validation-recall target gives about {high.false_alerts_per_1000:.0f} false alerts and {high.missed_r_per_1000:.0f} misses. Its actual test recall is {high.recall_r:.1%}, while only {high.precision_r:.1%} of alerts are correct.",
            "Each new cutoff was chosen on validation and then held fixed for test. A pooled recall target does not ensure that each drug reaches that target. These counts describe classification errors per candidate row, not per patient. They are not treatment rules or estimates of clinical benefit.",
        ],
        "07_label_coverage": [
            f"Among included E. coli events, {cover.loc[('escherichia coli', 'cefepime'), 'observed_panel_coverage']:.1%} have a retained cefepime label, compared with {cover.loc[('escherichia coli', 'gentamicin'), 'observed_panel_coverage']:.1%} for gentamicin. The six organisms were chosen by training frequency, rather than by test performance.",
            "The denominator contains organism-events with at least one clear, retained susceptible or resistant result in the chosen panel. This is coverage within the analytical data, not the raw testing rate. A blank or zero does not mean susceptibility or resistance. The data cannot establish why an outcome is absent or how the model would perform on unobserved combinations.",
        ],
        "08_missingness": [
            f"In test data, white blood cell count is missing for {missing.loc['lab_median_wbc_14d', 'missing_fraction']:.1%} of cultures. Prior candidate-resistance history is missing for {missing.loc['prior_candidate_resistance_days_ago', 'missing_fraction']:.1%} of culture-candidate pairs. A missing history record cannot establish that a patient had no past exposure or resistance.",
            "Shared clinical fields are counted once per culture. Candidate resistance history is counted once per culture and candidate. This avoids giving a culture extra weight just because it has more organism or drug rows. Similar missing rates across splits do not prove similar case mix or complete source records.",
        ],
    }
    highlights = [
        ("Prior history carries useful added information.", f"Validation resistance AP is {ap('core_only'):.3f} with core fields, {ap('core_prior_history'):.3f} with prior history, and {ap('full'):.3f} with all features. The paired comparison shows no clear extra gain from labs and vital signs under this recipe."),
        ("One cutoff behaves very differently across drugs.", f"It misses all {int(pip.n_r)} resistant piperacillin/tazobactam rows and flags all ampicillin rows. Good overall scores do not remove this problem."),
        ("Calibration reduces probability error.", f"Test Brier falls from {brier.raw:.3f} to {brier.calibrated:.3f}. Drug-level checks still matter, especially where there are few resistant patients."),
        ("More recall brings more false alerts.", f"Moving from the original cutoff to the 90% validation-recall target reduces misses from {original.missed_r_per_1000:.0f} to {high.missed_r_per_1000:.0f} per 1,000 rows, while false alerts rise from {original.false_alerts_per_1000:.0f} to {high.false_alerts_per_1000:.0f}."),
        ("Missing values and missing labels limit the claims.", "Availability patterns contain useful information, including recorded prior resistance. Missing outcomes remain unknown. Label coverage is conditional on inclusion in the analytical cohort."),
    ]
    methods = [
        ("Status and scope", "This appendix is an exploratory follow-up after the original test results were seen. It sits outside the original five-page report body. The original selected weighted LightGBM model, its predictions and its primary test scores remain fixed. No new winning model was selected."),
        ("Feature comparisons", "Four reduced models use the original proper-fit patients, held-out training calibration patients, balanced training weights and fixed LightGBM settings. Preprocessing is fitted only on proper-fit data. Reduced models are scored on validation only. Those validation patients had already informed the original model choice, so these are not fresh independent estimates."),
        ("Uncertainty", "New intervals use 500 shared patient-bootstrap draws. All rows for a sampled patient stay together. Paired comparisons use the same draws on both sides. The original model-comparison figure retains its 200-draw intervals. Intervals are pointwise, not adjusted for all comparisons, and do not include model-training or earlier selection uncertainty. Boundary intervals can collapse even when future performance is uncertain."),
        ("Calibration and cutoffs", "Probabilities refer to resistance, P(R) = 1 - P(S). Calibration was fitted on training patients. Reliability-bin boundaries and new recall-target cutoffs use validation predictions. They are then kept fixed for test. No test-fitted calibrator or threshold is introduced."),
        ("Units and labels", "The prediction unit is one patient, culture, organism and candidate drug. Labels are susceptible = 1 and resistant = 0. Intermediate, unknown and conflicting labels stay excluded. Error counts refer to candidate rows. Coverage refers only to organism-events with at least one retained panel label. Missingness uses culture or culture-candidate counts as stated in each table."),
        ("Limits that still apply", "The nine-drug list remains provisional until the instructor confirms it. The supplied tables do not include a hospital identifier, so hospital validation is unavailable. Shifted dates do not support a calendar trend claim. No feature comparison establishes causality, no threshold establishes treatment benefit, and no result establishes performance at a new site."),
        ("Source and reproduction", "The source is ARMD, described at https://www.nature.com/articles/s41597-025-05649-7 and released at https://datadryad.org/dataset/doi:10.5061/dryad.jq2bvq8kp. The gallery links each figure to its source tables. The workbook holds the result tables and metadata. scripts/start_advanced_analysis.py records the analysis plan and protects the primary files. scripts/make_advanced_package.py rebuilds this appendix and gallery from saved results."),
    ]
    for entry in index:
        if entry["stem"] not in notes:
            raise ValueError(f"No reviewed narrative for figure {entry['stem']}")
    for text in [*notes.values(), *highlights, *methods]:
        plain(text)
    return {"notes": notes, "highlights": highlights, "methods": methods,
            "test_rows": metadata["test_rows"], "test_patients": metadata["test_patients"],
            "validation_rows": ablation_metadata["partitions"]["validation"]["rows"],
            "validation_patients": ablation_metadata["partitions"]["validation"]["patients"]}


def draw_paragraph(pdf, text, x, top, width, size=12, leading=18, color=INK):
    style = ParagraphStyle("body", fontName="Helvetica", fontSize=size, leading=leading,
                           textColor=colors.HexColor("#" + color))
    paragraph = Paragraph(plain(text), style)
    _, height = paragraph.wrap(width, 10000)
    paragraph.drawOn(pdf, x, top - height)
    return top - height


def page_base(pdf, width, height, label, page, total):
    pdf.setPageSize((width, height))
    pdf.setFillColor(colors.HexColor("#" + TEAL))
    pdf.rect(0, height - 8, width, 8, stroke=0, fill=1)
    pdf.setFont("Helvetica-Bold", 10)
    pdf.setFillColor(colors.HexColor("#" + TEAL))
    pdf.drawString(44, height - 38, "ARMD  /  ADVANCED ANALYSIS")
    pdf.setFillColor(colors.HexColor("#" + GRAY))
    pdf.setFont("Helvetica", 10)
    pdf.drawRightString(width - 44, height - 38, label)
    pdf.setStrokeColor(colors.HexColor("#D8E3E8"))
    pdf.line(44, 35, width - 44, 35)
    pdf.setFont("Helvetica", 9)
    pdf.drawString(44, 20, "Exploratory appendix | Original selected model and primary results preserved")
    pdf.drawRightString(width - 44, 20, f"{page} / {total}")


def write_pdf(root, index, content):
    report = root / "report"
    report.mkdir(exist_ok=True)
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=(1008, 792))
    pdf.setTitle("ARMD: Advanced analysis appendix")
    pdf.setAuthor("ARMD teaching project")
    placements = []
    total = len(index) + 2
    page_base(pdf, 1008, 792, "How to read the extra analysis", 1, total)
    top = draw_paragraph(pdf, "Advanced analysis<br/>Antimicrobial resistance prediction", 44, 722, 900, 28, 33, NAVY)
    draw_paragraph(pdf, "Advanced analysis appendix | Eight figures with linked result tables", 44, top - 14, 900, 13, 18, GRAY)
    cards = [("Validation rows", f"{content['validation_rows']:,}", f"{content['validation_patients']:,} patients"),
             ("Test rows", f"{content['test_rows']:,}", f"{content['test_patients']:,} patients"),
             ("New uncertainty checks", "500 draws", "Patient bootstrap; model comparison uses 200")]
    for column, (label, value, detail) in enumerate(cards):
        x = 44 + column * 313
        pdf.setFillColor(colors.HexColor("#" + PALE))
        pdf.roundRect(x, 508, 296, 90, 10, fill=1, stroke=0)
        draw_paragraph(pdf, label, x + 16, 582, 267, 10, 13, GRAY)
        draw_paragraph(pdf, value, x + 16, 564, 267, 24, 27, NAVY)
        draw_paragraph(pdf, detail, x + 16, 530, 267, 9, 12, GRAY)
    top = 478
    for title, text in content["highlights"]:
        top = draw_paragraph(pdf, f"<b>{html.escape(title)}</b> {html.escape(text)}", 44, top, 920, 12, 17) - 13
    draw_paragraph(pdf, "Reading guide: 01 models; 02 feature groups; 03 drug performance; 04 overall calibration; 05 calibration by drug; 06 cutoff tradeoffs; 07 label coverage; 08 missingness.", 44, 134, 920, 11, 15, GRAY)
    draw_paragraph(pdf, "The feature comparisons use validation only. The other model checks describe the already examined test set. Treat every new finding as exploratory.", 44, 87, 920, 10, 14, TEAL)
    pdf.showPage()
    for position, entry in enumerate(index, start=1):
        width, height = (792, 1008) if entry["stem"] == "05_calibration_by_drug" else (1008, 792)
        page_base(pdf, width, height, f"Figure {position:02d} of {len(index):02d}", position + 1, total)
        placements.append((position, root / "results" / "advanced" / entry["files"]["png"],
                           fitz.Rect(28, 65, width - 28, height - 211)))
        top = 195
        for paragraph in content["notes"][entry["stem"]]:
            top = draw_paragraph(pdf, html.escape(paragraph), 44, top, width - 88, 11, 15) - 9
        if top < 43:
            raise ValueError(f"Narrative overflows figure page: {entry['stem']}")
        pdf.showPage()
    page_base(pdf, 1008, 792, "Methods, limits and source files", total, total)
    top = draw_paragraph(pdf, "What these results can support", 44, 723, 920, 25, 31, NAVY) - 23
    for title, paragraph in content["methods"]:
        top = draw_paragraph(pdf, f"<b>{html.escape(title)}.</b> {html.escape(paragraph)}", 44, top, 920, 11, 16) - 15
    if top < 44:
        raise ValueError("Methods page exceeds its available space.")
    pdf.showPage()
    pdf.save()
    document = fitz.open(stream=buffer.getvalue(), filetype="pdf")
    for page_number, source, rect in placements:
        document[page_number].insert_image(rect, filename=str(source), keep_proportion=True)
    document.set_toc([[1, "What the extra analysis adds", 1]] +
                     [[1, f"{i + 1:02d}. {entry['title']}", i + 2] for i, entry in enumerate(index)] +
                     [[1, "Methods and limits", total]])
    all_text = "\n".join(page.get_text() for page in document)
    plain(all_text)
    if len(document) != total:
        raise AssertionError("Unexpected appendix page count.")
    destination = report / "advanced_analysis.pdf"
    document.save(destination, garbage=4, deflate=True)
    document.close()
    return total


def write_gallery(root, index, content):
    esc = html.escape
    nav = "".join(f'<a href="#{esc(entry["stem"])}"><span>{number:02d}</span>{esc(entry["title"])}</a>'
                  for number, entry in enumerate(index, start=1))
    highlights = "".join(f'<li><strong>{esc(title)}</strong> {esc(text)}</li>' for title, text in content["highlights"])
    articles = []
    for number, entry in enumerate(index, start=1):
        sources = " ".join(f'<a href="../results/advanced/{esc(path)}">{esc(Path(path).name)}</a>' for path in entry["source_tables"])
        links = f'<a href="../results/advanced/{esc(entry["files"]["png"])}" target="_blank" rel="noopener">PNG</a>'
        paragraphs = "".join(f'<p>{esc(paragraph)}</p>' for paragraph in content["notes"][entry["stem"]])
        articles.append(f'''<article id="{esc(entry['stem'])}"><div class="eyebrow">FIGURE {number:02d} / {len(index):02d}</div>
<h2>{esc(entry['title'])}</h2><p class="question">{esc(entry['question'])}</p>
<a class="image-link" href="../results/advanced/{esc(entry['files']['png'])}" target="_blank" rel="noopener" aria-label="Open full-size figure {number}">
<img loading="lazy" src="../results/advanced/{esc(entry['files']['png'])}" alt="{esc(entry['title'])}. {esc(content['notes'][entry['stem']][0])}"></a>
<div class="interpretation">{paragraphs}</div><div class="files"><b>Export figure</b>{links}<a href="#top">Back to top</a></div>
<details><summary>Source tables and scope</summary><p>{esc(entry['note'])}</p><div class="source-links">{sources}</div></details></article>''')
    methods = "".join(f'<h3>{esc(title)}</h3><p>{esc(text)}</p>' for title, text in content["methods"])
    document = f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>ARMD | Advanced analysis</title><style>
:root{{--ink:#233e50;--navy:#17354a;--teal:#117c83;--muted:#526877;--line:#d8e3e8;--pale:#edf5f5;--paper:#fff}}*{{box-sizing:border-box}}html{{scroll-behavior:smooth;scroll-padding-top:20px}}body{{margin:0;background:#f4f7f9;color:var(--ink);font:16px/1.65 system-ui,-apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif}}a{{color:var(--teal);text-underline-offset:3px}}a:hover{{color:var(--navy)}}a:focus-visible{{outline:3px solid #ca762e;outline-offset:5px}}aside{{position:fixed;top:0;left:0;bottom:0;width:280px;padding:28px 22px;background:var(--paper);border-right:1px solid var(--line);overflow:auto}}.brand{{color:var(--teal);font-size:12px;font-weight:750;letter-spacing:.15em}}aside h2{{font-size:22px;line-height:1.3;margin:10px 0 20px}}nav a{{display:block;text-decoration:none;font-size:13px;line-height:1.5;padding:10px 0;color:var(--muted);border-bottom:1px solid #edf1f3}}nav span{{display:block;color:var(--teal);font-size:11px;font-weight:700;letter-spacing:.07em}}main{{margin-left:280px;padding:38px clamp(22px,4vw,64px);max-width:1680px}}header,article,.methods{{background:var(--paper);border:1px solid var(--line);border-radius:16px;padding:clamp(22px,3vw,40px);margin-bottom:28px;box-shadow:0 4px 24px #17354a04}}header{{border-top:6px solid var(--teal)}}h1{{font-size:clamp(32px,3.3vw,48px);line-height:1.12;letter-spacing:-.025em;color:var(--navy);margin:12px 0 20px;max-width:800px}}h2{{font-size:clamp(23px,2.1vw,31px);line-height:1.25;color:var(--navy);margin:10px 0 12px}}h3{{font-size:17px;color:var(--navy);margin:25px 0 5px}}.eyebrow{{font-size:12px;font-weight:750;letter-spacing:.11em;color:var(--teal)}}.lead,.question{{color:var(--muted);max-width:900px}}.stats{{display:grid;grid-template-columns:repeat(3,1fr);gap:12px;margin:24px 0}}.stat{{background:var(--pale);padding:16px;border-radius:10px;font-size:12px;line-height:1.5;color:var(--muted)}}.stat b{{display:block;font-size:27px;color:var(--navy)}}.findings{{padding-left:22px}}.findings li{{padding-left:5px;margin:12px 0}}.notice{{padding:14px 18px;border-left:3px solid var(--teal);background:var(--pale);font-size:14px}}.downloads,.files{{display:flex;gap:14px;align-items:center;flex-wrap:wrap;margin-top:20px;font-size:13px}}.downloads a{{padding:9px 13px;border:1px solid var(--line);border-radius:7px;text-decoration:none}}.image-link{{display:block;margin:20px -12px 14px}}img{{display:block;width:100%;height:auto}}.interpretation{{max-width:1040px}}.interpretation p{{margin:13px 0}}details{{border-top:1px solid var(--line);margin-top:18px;padding-top:14px;font-size:13px;color:var(--muted)}}summary{{cursor:pointer;font-weight:600}}.source-links a{{display:inline-block;margin:0 14px 8px 0}}footer{{font-size:12px;color:var(--muted);padding:5px 0 24px}}@media(max-width:1000px){{aside{{position:static;width:auto;border-right:0;border-bottom:1px solid var(--line)}}nav{{display:grid;grid-template-columns:repeat(2,1fr);gap:0 22px}}main{{margin:0;padding:20px}}.stats{{grid-template-columns:1fr}}}}@media print{{aside,.downloads,.files,details{{display:none}}main{{margin:0;padding:0}}header,article,.methods{{break-inside:avoid;box-shadow:none;border:0}}article{{page-break-before:always}}body{{background:white;font-size:12px}}}}
</style></head><body id="top"><aside><div class="brand">ARMD / RESEARCH PROJECT</div><h2>Advanced analysis</h2><nav><a href="#overview">Overview</a>{nav}<a href="#methods">Methods and limits</a></nav></aside>
<main><header id="overview"><div class="eyebrow">EXPLORATORY APPENDIX</div><h1>Advanced analysis<br>Antimicrobial resistance prediction</h1><p class="lead">Eight figures examine feature value, drug-level errors, probability quality and gaps in the data. Each figure has a short explanation and links to the numbers behind it.</p>
<div class="stats"><div class="stat">Validation rows<b>{content['validation_rows']:,}</b>{content['validation_patients']:,} patients</div><div class="stat">Test rows<b>{content['test_rows']:,}</b>{content['test_patients']:,} patients</div><div class="stat">New uncertainty checks<b>500 draws</b>Patient bootstrap; original model figure uses 200</div></div>
<ul class="findings">{highlights}</ul><p class="notice">These are exploratory follow-up analyses. The original selected model and primary results stay fixed. Feature comparisons use validation only; the test set was already examined in the original report.</p>
<div class="downloads"><a href="advanced_analysis.pdf">Download the 10-page appendix</a><a href="../results/advanced/advanced_results.xlsx">Open the results workbook</a><a href="report.pdf">Main five-page report and references</a></div></header>
{''.join(articles)}<section class="methods" id="methods"><div class="eyebrow">HOW TO READ THE EVIDENCE</div><h2>Methods and limits</h2>{methods}</section>
<footer>Charts are available as PNG. This gallery uses local files and needs no web connection. Keep it inside the project folder so its figure and table links work.</footer></main></body></html>'''
    plain(document)
    (root / "report" / "advanced_analysis.html").write_text(document, encoding="utf-8")


def flatten(prefix, value):
    if isinstance(value, dict):
        for key, child in value.items():
            yield from flatten(f"{prefix}.{key}" if prefix else key, child)
    elif isinstance(value, list):
        for number, child in enumerate(value):
            yield from flatten(f"{prefix}[{number}]", child)
    else:
        yield prefix, value


def style_sheet(sheet, table_name):
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 42
    for column in sheet.columns:
        sample = list(column)[:200]
        max_length = max(len(str(cell.value or "")) for cell in sample)
        sheet.column_dimensions[column[0].column_letter].width = min(48, max(13, max_length + 2))
    if sheet.max_row > 1:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False,
                                               showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
    sheet.auto_filter.ref = sheet.dimensions


def write_workbook(root, index):
    folder = root / "results" / "advanced"
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Index"
    overview.append(["Sheet", "Rows", "Purpose", "Source file"])
    for number, source in enumerate(sorted(folder.glob("*.csv")), start=1):
        if "bootstrap" in source.stem:
            continue
        frame = pd.read_csv(source)
        name = source.stem[:31]
        if name in workbook.sheetnames:
            raise ValueError(f"Truncated workbook sheet name collision: {name}")
        sheet = workbook.create_sheet(name)
        sheet.append(frame.columns.tolist())
        for row in frame.itertuples(index=False, name=None):
            converted = []
            for value in row:
                if pd.isna(value):
                    converted.append(None)
                elif hasattr(value, "item"):
                    converted.append(value.item())
                else:
                    converted.append(value)
            sheet.append(converted)
        for column in sheet.iter_cols(min_row=2):
            for cell in column:
                if isinstance(cell.value, float) and math.isfinite(cell.value):
                    cell.number_format = "0.0000" if abs(cell.value) < 2 else "#,##0.00"
                elif isinstance(cell.value, int) and not isinstance(cell.value, bool):
                    cell.number_format = "#,##0"
        style_sheet(sheet, f"Results{number:02d}")
        overview.append([name, len(frame), SHEET_NOTES.get(source.stem, "Advanced analysis result table."), source.name])
        overview.cell(overview.max_row, 1).hyperlink = f"#'{name}'!A1"
        overview.cell(overview.max_row, 1).style = "Hyperlink"
    metadata = workbook.create_sheet("Metadata")
    metadata.append(["Source", "Field", "Value"])
    for source in ["analysis_protocol.json", "ablation_metadata.json", "diagnostics_metadata.json",
                   "fixed_thresholds_validation.json", "fixed_calibration_bins_validation.json"]:
        for key, value in flatten("", load_json(folder / source)):
            metadata.append([source, key, value])
    style_sheet(metadata, "MetadataTable")
    metadata.column_dimensions["A"].width = 37
    metadata.column_dimensions["B"].width = 58
    metadata.column_dimensions["C"].width = 100
    for row in metadata.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
        if isinstance(row[2].value, str) and len(row[2].value) > 95:
            metadata.row_dimensions[row[2].row].height = min(90, 15 * math.ceil(len(row[2].value) / 95))
    figures = workbook.create_sheet("Figure index")
    figures.append(["Figure", "Title", "Question", "Source tables", "Scope note", "PNG file"])
    for entry in index:
        figures.append([entry["stem"], entry["title"], entry["question"], "; ".join(entry["source_tables"]),
                        entry["note"], entry["files"]["png"]])
    style_sheet(figures, "FigureIndexTable")
    overview.append(["Metadata", metadata.max_row - 1, "Study scope, fixed model settings, bootstrap details and source checksums.", "Saved JSON metadata"])
    overview.append(["Figure index", len(index), "Figure questions, source tables and export paths.", "figure_index.json"])
    overview.append(["Interpretation", None, "Exploratory follow-up. Fixed test model. Ablations are validation only. Missing cells are not zeros. Intervals are pointwise and conditional on fitted models.", None])
    overview.append(["Bootstrap draws", None, "Raw replicate CSV files are kept separately; they are not duplicated in this workbook.", "ablation_bootstrap_replicates.csv"])
    style_sheet(overview, "WorkbookIndexTable")
    overview.column_dimensions["A"].width = 34
    overview.column_dimensions["B"].width = 14
    overview.column_dimensions["C"].width = 105
    overview.column_dimensions["D"].width = 43
    for row in overview.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
        overview.row_dimensions[row[2].row].height = 34
    workbook.properties.title = "ARMD advanced analysis results"
    workbook.properties.subject = "Exploratory analyses with the original selected model held fixed"
    workbook.save(folder / "advanced_results.xlsx")
    return len(workbook.sheetnames)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--include-reports", action="store_true", help="Also create the local PDF appendix and HTML gallery.")
    args = parser.parse_args()
    root = args.root.resolve()
    folder = root / "results" / "advanced"
    index = load_json(folder / "figure_index.json")
    if len(index) != 8 or len({entry["stem"] for entry in index}) != 8:
        raise ValueError("Expected the eight reviewed advanced figures.")
    for entry in index:
        for path in list(entry["files"].values()) + entry["source_tables"]:
            resolved = (folder / path).resolve()
            if not resolved.is_relative_to(root) or not resolved.is_file():
                raise ValueError(f"Missing or out-of-project figure resource: {path}")
    pages = None
    outputs = ["results/advanced/advanced_results.xlsx"]
    if args.include_reports:
        content = build_content(root, index)
        pages = write_pdf(root, index, content)
        write_gallery(root, index, content)
        outputs += ["report/advanced_analysis.pdf", "report/advanced_analysis.html"]
    sheets = write_workbook(root, index)
    print(json.dumps({"appendix_pages": pages, "figures": len(index), "workbook_sheets": sheets,
                      "outputs": outputs}, indent=2))


if __name__ == "__main__":
    main()
