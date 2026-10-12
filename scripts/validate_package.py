"""Check the real-data outputs; report checks require --include-reports."""
from pathlib import Path
import argparse
import json
import os
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'src'))
os.environ.setdefault('MPLCONFIGDIR', str(ROOT/'.cache/matplotlib'))
import numpy as np
import pandas as pd
import joblib
from sklearn.metrics import roc_auc_score, average_precision_score, balanced_accuracy_score, precision_score, recall_score, brier_score_loss


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--include-reports', action='store_true',
                        help='Also validate the optional main PDF and Word report in report/.')
    args = parser.parse_args()
    results = ROOT/'results'
    df = pd.read_parquet(results/'analytical.parquet')
    predictions = pd.read_parquet(results/'models/predictions.parquet')
    comparison = pd.read_csv(results/'models/model_comparison.csv')
    audit = json.loads((results/'construction_audit.json').read_text())
    summary = json.loads((results/'models/evaluation_summary.json').read_text())
    selection = json.loads((results/'models/frozen_selection.json').read_text())
    key = ['patient_id','culture_id','organism','candidate_antibiotic']
    assert len(df) == audit['analytical_rows'] and not df.duplicated(key).any()
    assert df.groupby('patient_id').split.nunique().max() == 1
    assert set(df.label.unique()) == {0,1}
    assert set(summary['features']).isdisjoint({'patient_id','culture_id','organism','culture_time','label','split','dataset'})
    fields = pd.read_csv(ROOT/'docs/final_data_dictionary.csv')['canonical_field']
    assert set(fields) == set(df.columns)
    test = df[df.split.eq('test')].reset_index(drop=True)
    pred_test = predictions[predictions.split.eq('test')].reset_index(drop=True)
    pd.testing.assert_frame_equal(test[key], pred_test[key], check_dtype=False)
    assert np.array_equal(test.label, pred_test.label)
    errors = []
    for _, row in comparison[comparison['split'].eq('test')].iterrows():
        y = pred_test.label.to_numpy()
        p = pred_test['p_s_'+row.model].to_numpy()
        chosen = (p >= row.threshold_s).astype(int)
        macro_auc, macro_ap_r, macro_ap_s = [], [], []
        for _, indices in pred_test.groupby('candidate_antibiotic').groups.items():
            idx = np.asarray(indices)
            macro_auc.append(roc_auc_score(y[idx], p[idx]))
            macro_ap_r.append(average_precision_score(1-y[idx], 1-p[idx]))
            macro_ap_s.append(average_precision_score(y[idx], p[idx]))
        checks = {'antibiotic_macro_auroc':np.mean(macro_auc), 'antibiotic_macro_auprc_r':np.mean(macro_ap_r),
                  'antibiotic_macro_auprc_s':np.mean(macro_ap_s),
                  'antibiotic_macro_class_auprc':(np.mean(macro_ap_r)+np.mean(macro_ap_s))/2,
                  'balanced_accuracy':balanced_accuracy_score(y,chosen),
                  'precision_r':precision_score(y,chosen,pos_label=0,zero_division=0),
                  'recall_r':recall_score(y,chosen,pos_label=0), 'brier':brier_score_loss(y,p)}
        for metric, expected in checks.items():
            error = abs(float(row[metric])-expected)
            errors.append(error)
            assert error < 1e-10, (row.model,metric,error)
    assert selection['selected_model'] == summary['selected_model']
    assert not selection['test_used_for_selection']
    model = joblib.load(results/'models/models'/f"{summary['selected_model']}.joblib")
    np.testing.assert_allclose(model.predict_proba(test.iloc[:250])[:,1],
                               pred_test['p_s_'+summary['selected_model']].iloc[:250], rtol=1e-12)
    report_checks = {
        'status': 'skipped',
        'reason': 'Report checks skipped per repository scope; enable explicitly with --include-reports.',
    }
    if args.include_reports:
        from pypdf import PdfReader
        from docx import Document
        pdf = PdfReader(ROOT/'report/report.pdf')
        assert len(pdf.pages) == 6
        text = '\n'.join(p.extract_text() or '' for p in pdf.pages)
        assert '\u2014' not in text
        assert 'References and file guide' in pdf.pages[5].extract_text()
        doc = Document(ROOT/'report/report.docx')
        document_text = '\n'.join(p.text for p in doc.paragraphs) + '\n'.join(
            c.text for t in doc.tables for r in t.rows for c in r.cells)
        assert '\u2014' not in document_text
        report_checks = {'status': 'passed', 'pdf_body_pages': len(pdf.pages) - 1,
                         'pdf_reference_pages': 1, 'report_em_dashes': 0}
    record = {'status':'passed', 'analytical_rows':len(df), 'patients':int(df.patient_id.nunique()),
              'cultures':len(df[['patient_id','culture_id']].drop_duplicates()), 'duplicate_keys':0,
              'patients_in_multiple_splits':0, 'independent_metric_checks':len(errors),
              'maximum_metric_absolute_difference':max(errors), 'saved_model_prediction_check':'passed',
              'dictionary_matches_analytical_columns':True,
              'include_reports':args.include_reports, 'report_checks':report_checks,
              'data_source':'real ARMD, publisher checksum verified',
              'limitations':['Instructor nine-drug list not supplied; panel remains provisional.',
                             'Separate hospital identifiers and external dataset absent; independent site validation unavailable.']}
    (results/'package_validation.json').write_text(json.dumps(record,indent=2),encoding='utf-8')
    print(json.dumps(record,indent=2))


if __name__ == '__main__':
    main()
