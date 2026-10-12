"""Check advanced results; reports are excluded unless --include-reports is set.

--require-present requires the workbook and other repository deliverables. It
requires the appendix and gallery too when combined with --include-reports.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import re
import sys

import numpy as np
import pandas as pd
from PIL import Image
from pypdf import PdfReader
from sklearn.metrics import average_precision_score, brier_score_loss, precision_score, recall_score, roc_auc_score

ROOT = Path(__file__).resolve().parents[1]
KEY = ["patient_id", "culture_id", "organism", "candidate_antibiotic"]


def validate(root: Path, require_present: bool = False, include_reports: bool = False) -> dict:
    root = root.resolve()
    advanced, models = root / "results/advanced", root / "results/models"
    record = {"status": "running", "require_present": require_present,
              "include_reports": include_reports, "sections": {},
              "limitations": ["Numerical/file integrity checks do not establish clinical usefulness or external validity.",
                              "The advanced test analyses are exploratory because the original test results were already reviewed."]}
    comparisons = []

    def close(actual, expected, label):
        if pd.isna(actual) and pd.isna(expected):
            return
        error = abs(float(actual) - float(expected))
        if not np.isfinite(error) or error > 1e-10:
            raise AssertionError(f"{label}: found {actual}, expected {expected}")
        comparisons.append(error)

    def read(name):
        return pd.read_csv(advanced / name, float_precision="round_trip")

    def sha(path):
        digest = hashlib.sha256()
        with path.open("rb") as stream:
            for block in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(block)
        return digest.hexdigest()

    try:
        protocol = json.loads((advanced / "analysis_protocol.json").read_text(encoding="utf-8"))
        protected = protocol["protected_sha256"]
        assert len(protected) == 7, "Expected the seven original protected artifacts."
        for name, expected in protected.items():
            assert sha(root / name) == expected, f"Protected original artifact changed: {name}"
        record["sections"]["original_artifacts"] = {"unchanged_sha256": len(protected)}

        data = pd.read_parquet(root / "results/analytical.parquet")
        predictions = pd.read_parquet(models / "predictions.parquet")
        selection = json.loads((models / "frozen_selection.json").read_text(encoding="utf-8"))
        model = selection["selected_model"]
        threshold_s = selection["thresholds_s"][model]
        pcol, rawcol = "p_s_" + model, "p_s_raw_" + model
        assert selection["test_used_for_selection"] is False
        assert selection["selection_split"] == "validation"
        assert not data.duplicated(KEY).any() and not predictions.duplicated(KEY).any()
        assert data.groupby("patient_id", observed=True).split.nunique().max() == 1
        expected = data[data.split.isin(["validation", "test"])][KEY + ["label", "split"]]
        aligned = expected.merge(predictions[KEY + ["label", "split"]], on=KEY, how="outer",
                                 validate="one_to_one", suffixes=("_data", "_pred"), indicator=True)
        assert aligned._merge.eq("both").all()
        assert aligned.label_data.eq(aligned.label_pred).all() and aligned.split_data.eq(aligned.split_pred).all()
        partitions = {split: predictions[predictions.split == split].reset_index(drop=True)
                      for split in ["validation", "test"]}
        test, validation = partitions["test"], partitions["validation"]
        record["sections"]["source_alignment"] = {"prediction_rows": len(predictions),
            "duplicate_keys": 0, "patients_in_multiple_splits": 0, "label_or_split_mismatches": 0}

        drug_table = read("drug_performance.csv")
        panel = json.loads((root / "config.json").read_text(encoding="utf-8"))["antibiotics"]
        assert len(drug_table) == len(panel) == 9 and set(drug_table.candidate_antibiotic) == set(panel)
        for row in drug_table.itertuples(index=False):
            group = test[test.candidate_antibiotic == row.candidate_antibiotic]
            y, p = group.label.to_numpy(int), group[pcol].to_numpy(float)
            flags, r = p < threshold_s, y == 0
            predicted = (~flags).astype(int)
            checks = {"n": len(group), "n_r": r.sum(), "n_s": y.sum(), "patients": group.patient_id.nunique(),
                "patients_with_r": group.loc[r, "patient_id"].nunique(),
                "patients_with_s": group.loc[~r, "patient_id"].nunique(),
                "n_flagged": flags.sum(), "false_alerts": ((~r) & flags).sum(), "missed_r": (r & ~flags).sum(),
                "auroc": roc_auc_score(y, p), "ap_r": average_precision_score(1 - y, 1 - p),
                "brier": brier_score_loss(y, p), "recall_r": recall_score(y, predicted, pos_label=0),
                "precision_r": precision_score(y, predicted, pos_label=0, zero_division=np.nan),
                "prevalence_r": r.mean(), "flagged_fraction": flags.mean(), "threshold_s": threshold_s}
            checks["ap_lift"] = checks["ap_r"] / checks["prevalence_r"]
            checks["ap_above_chance"] = (checks["ap_r"] - checks["prevalence_r"]) / (1 - checks["prevalence_r"])
            for metric, expected in checks.items():
                close(getattr(row, metric), expected, f"{row.candidate_antibiotic}/{metric}")
        record["sections"]["drug_performance"] = {"drugs": len(drug_table), "independent_library": "scikit-learn"}

        fixed = json.loads((advanced / "fixed_thresholds_validation.json").read_text(encoding="utf-8"))
        assert fixed["test_used_for_threshold_selection"] is False and fixed["selection_split"] == "validation"
        assert fixed["selected_model"] == model
        assert [p["operating_point"] for p in fixed["operating_points"]] == ["original", "recall_70", "recall_80", "recall_90"]
        assert [p["target_recall_r"] for p in fixed["operating_points"]] == [None, .7, .8, .9]
        operating = read("threshold_operating_points.csv")
        assert len(operating) == 8
        for point in fixed["operating_points"]:
            if point["operating_point"] == "original":
                assert point["threshold_s"] == threshold_s
            else:
                # Independent order statistic of resistant validation scores.
                scores_r = np.sort(validation.loc[validation.label == 0, pcol].to_numpy(float))
                rank = int(np.ceil(point["target_recall_r"] * len(scores_r))) - 1
                expected_threshold = np.nextafter(scores_r[rank], np.inf)
                assert point["threshold_s"] == expected_threshold, "Recall cutoff differs from validation resistant order statistic."
                recall = np.mean(scores_r < point["threshold_s"])
                stricter_recall = np.mean(scores_r < np.nextafter(point["threshold_s"], -np.inf))
                assert recall >= point["target_recall_r"] > stricter_recall
            for split, frame in partitions.items():
                rows = operating[(operating.operating_point == point["operating_point"]) & (operating.split == split)]
                assert len(rows) == 1
                row = rows.iloc[0]
                r = frame.label.to_numpy(int) == 0
                flagged = frame[pcol].to_numpy(float) < point["threshold_s"]
                tp, fp, fn = (r & flagged).sum(), (~r & flagged).sum(), (r & ~flagged).sum()
                checks = {"threshold_s": point["threshold_s"], "threshold_r": 1 - point["threshold_s"],
                    "n": len(frame), "n_r": r.sum(), "n_flagged": flagged.sum(), "true_r": tp,
                    "false_alerts": fp, "missed_r": fn, "patients": frame.patient_id.nunique(),
                    "recall_r": tp / r.sum(), "precision_r": tp / flagged.sum() if flagged.sum() else np.nan,
                    "flagged_fraction": flagged.mean(), "false_alerts_per_1000": 1000 * fp / len(frame),
                    "missed_r_per_1000": 1000 * fn / len(frame)}
                for name, expected in checks.items():
                    close(row[name], expected, f"{split}/{point['operating_point']}/{name}")
        curve = read("validation_tradeoff_curve.csv")
        validation_r = validation.label.to_numpy(int) == 0
        scores = validation[pcol].to_numpy(float)
        unique_scores, inverse, count = np.unique(scores, return_inverse=True, return_counts=True)
        count_r = np.bincount(inverse, weights=validation_r, minlength=len(unique_scores))
        expected_flagged, expected_tp = np.r_[0, count.cumsum()], np.r_[0, count_r.cumsum()]
        assert np.array_equal(curve.threshold_s.to_numpy(), np.r_[0., np.nextafter(unique_scores, np.inf)])
        assert np.array_equal(curve.n_flagged.to_numpy(), expected_flagged)
        assert np.array_equal(curve.true_r.to_numpy(), expected_tp)
        np.testing.assert_allclose(curve.recall_r, expected_tp / validation_r.sum(), atol=1e-12)
        record["sections"]["thresholds"] = {"frozen_points": 4, "validated_split_point_rows": len(operating),
            "validation_recall_targets": [p["target_recall_r"] for p in fixed["operating_points"] if p["target_recall_r"] is not None],
            "unique_score_tradeoff_rows": len(curve), "tie_rule": "strict p_s < threshold_s verified"}

        bin_spec = json.loads((advanced / "fixed_calibration_bins_validation.json").read_text(encoding="utf-8"))
        assert bin_spec["selection_split"] == "validation" and not bin_spec["test_used_for_bin_selection"]
        bins, summaries, histogram = read("calibration_bins.csv"), read("calibration_summary.csv"), read("probability_histogram.csv")
        for spec in bin_spec["bins"]:
            column = pcol if spec["probability"] == "calibrated" else rawcol
            val_group = validation if spec["scope"] == "overall" else validation[validation.candidate_antibiotic == spec["candidate_antibiotic"]]
            test_group = test if spec["scope"] == "overall" else test[test.candidate_antibiotic == spec["candidate_antibiotic"]]
            internal = np.quantile(1 - val_group[column].to_numpy(float), np.arange(1, spec["requested_bins"]) / spec["requested_bins"])
            expected_edges = np.r_[0., np.unique(internal[(internal > 0) & (internal < 1)]), 1.]
            np.testing.assert_array_equal(spec["edges"], expected_edges)
            select = lambda df: df[(df.scope == spec["scope"]) & (df.candidate_antibiotic == spec["candidate_antibiotic"]) & (df.probability == spec["probability"])]
            p_r, r = 1 - test_group[column].to_numpy(float), 1 - test_group.label.to_numpy(int)
            codes = pd.cut(p_r, expected_edges, labels=False, right=False, include_lowest=True)
            codes[p_r == 1] = len(expected_edges) - 2
            chosen = select(bins).sort_values("bin")
            assert len(chosen) == len(expected_edges) - 1
            assert chosen.n.sum() == len(test_group) and chosen.n_r.sum() == r.sum()
            ece = 0
            for row in chosen.itertuples(index=False):
                mask = codes == row.bin - 1
                n = mask.sum()
                close(row.edge_left, expected_edges[row.bin - 1], "bin edge left")
                close(row.edge_right, expected_edges[row.bin], "bin edge right")
                close(row.n, n, "bin n")
                close(row.n_r, r[mask].sum(), "bin resistant n")
                close(row.patients, test_group.loc[mask, "patient_id"].nunique(), "bin patients")
                assert row.sparse == bool(n < 100 or r[mask].sum() < 20)
                for field, expected in {"mean_predicted_r": p_r[mask].mean() if n else np.nan,
                    "observed_r": r[mask].mean() if n else np.nan,
                    "calibration_gap": (p_r[mask] - r[mask]).mean() if n else np.nan}.items():
                    close(getattr(row, field), expected, field)
                ece += abs(np.sum(p_r[mask] - r[mask])) / len(test_group)
            summary = select(summaries).iloc[0]
            for field, expected in {"brier": brier_score_loss(r, p_r), "prevalence_r": r.mean(),
                "mean_predicted_r": p_r.mean(), "calibration_gap": (p_r - r).mean(), "ece_validation_bins": ece}.items():
                close(summary[field], expected, f"calibration summary/{field}")
            hist = select(histogram).sort_values("bin")
            counts, edges = np.histogram(p_r, np.linspace(0, 1, 21))
            np.testing.assert_array_equal(hist.n, counts)
        record["sections"]["calibration"] = {"validation_defined_groups": len(bin_spec["bins"]),
            "bins_checked": len(bins), "histogram_groups": len(bin_spec["bins"]), "test_fitted_calibration": False}
        calibration_comparison = read("calibration_comparison.csv")
        overall_calibration = summaries[summaries.scope == "overall"].set_index("probability")
        for row in calibration_comparison.itertuples(index=False):
            raw, calibrated = overall_calibration.loc["raw", row.metric], overall_calibration.loc["calibrated", row.metric]
            close(row.raw, raw, "calibration comparison raw")
            close(row.calibrated, calibrated, "calibration comparison calibrated")
            close(row.raw_minus_calibrated, raw - calibrated, "calibration comparison difference")

        ablation = pd.read_parquet(advanced / "ablation_predictions.parquet")
        ablation_meta = json.loads((advanced / "ablation_metadata.json").read_text(encoding="utf-8"))
        assert set(ablation.split) == {"validation"}
        assert ablation_meta["test_rows_loaded"] is False and ablation_meta["test_used_for_fitting_or_selection"] is False
        assert ablation_meta["hyperparameters_or_thresholds_tuned"] is False
        joined = validation[KEY + ["label", pcol]].merge(ablation, on=KEY, validate="one_to_one", how="outer", indicator=True)
        assert joined._merge.eq("both").all() and joined.label_x.eq(joined.label_y).all()
        np.testing.assert_array_equal(joined[pcol], joined.p_s_full)
        ablation_metrics, paired = read("ablation_metrics.csv"), read("ablation_paired_differences.csv")
        replicate = read("ablation_bootstrap_replicates.csv")
        assert set(ablation_metrics.split) == set(paired.split) == {"validation"}
        for model_name, group in ablation_metrics.groupby("model"):
            p = ablation["p_s_" + model_name].to_numpy(float)
            y = ablation.label.to_numpy(int)
            aucs, aps = [], []
            for _, idx in ablation.groupby("candidate_antibiotic", observed=True).indices.items():
                aucs.append(roc_auc_score(y[idx], p[idx]))
                aps.append(average_precision_score(1 - y[idx], 1 - p[idx]))
            expected = {"antibiotic_macro_auroc": np.mean(aucs), "antibiotic_macro_auprc_r": np.mean(aps),
                        "brier": brier_score_loss(y, p)}
            for row in group.itertuples(index=False):
                close(row.estimate, expected[row.metric], f"ablation/{model_name}/{row.metric}")
                values = replicate.loc[replicate.model == model_name, row.metric].dropna()
                low, high = np.quantile(values, [.025, .975])
                close(row.ci_low, low, "ablation CI low")
                close(row.ci_high, high, "ablation CI high")
                assert len(values) == row.valid_replicates == row.requested_replicates == 500
        for row in paired.itertuples(index=False):
            points = ablation_metrics[ablation_metrics.metric == row.metric].set_index("model").estimate
            close(row.difference_a_minus_b, points[row.model_a] - points[row.model_b], "paired difference")
            draws = replicate.pivot(index="replicate", columns="model", values=row.metric)
            low, high = np.quantile(draws[row.model_a] - draws[row.model_b], [.025, .975])
            close(row.ci_low, low, "paired CI low")
            close(row.ci_high, high, "paired CI high")
        record["sections"]["ablation"] = {"validation_rows": len(ablation), "feature_sets": ablation_metrics.model.nunique(),
            "metric_estimates": len(ablation_metrics), "paired_intervals": len(paired),
            "full_probabilities_exactly_equal_original": True, "test_rows_in_ablation": 0,
            "new_model_selection": False}

        missingness = read("missingness_by_split.csv")
        missing_features = json.loads((models / "evaluation_summary.json").read_text(encoding="utf-8"))["numeric_features"]
        assert set(missingness.feature) == set(missing_features)
        assert not missingness.duplicated(["feature", "split"]).any()
        unit_checks = []
        for feature in missing_features:
            keys = ["patient_id", "culture_id"] + (["candidate_antibiotic"] if feature.startswith("prior_candidate") else [])
            # Include NaN in uniqueness: a missing/observed disagreement must fail.
            value_counts = data.groupby(keys, observed=True, dropna=False)[feature].nunique(dropna=False)
            assert not (value_counts > 1).any(), f"Values disagree before unit deduplication: {feature}"
            assert data.groupby(keys, observed=True, dropna=False).split.nunique(dropna=False).max() == 1
            frame = data[keys + ["split", feature]].drop_duplicates(keys)
            for split, group in frame.groupby("split", observed=True):
                row = missingness[(missingness.feature == feature) & (missingness.split == split)].iloc[0]
                close(row.units, len(group), "missingness units")
                close(row.missing_n, group[feature].isna().sum(), "missingness n")
                close(row.missing_fraction, group[feature].isna().mean(), "missingness fraction")
                assert row.unit == ("culture-candidate" if feature.startswith("prior_candidate") else "culture")
            unit_checks.append({"feature": feature, "units_all_splits": len(frame), "conflicts_before_deduplication": 0})
        record["sections"]["missingness"] = {"rows": len(missingness), "unit_checks": unit_checks}

        slices = read("organism_drug_slices.csv")
        event = ["patient_id", "culture_id", "organism"]
        training_events = data.loc[data.split == "train", event].drop_duplicates()
        top = training_events.groupby("organism", observed=True).size().rename("count").reset_index().sort_values(["count", "organism"], ascending=[False, True]).head(6)
        assert set(slices.organism) == set(top.organism)
        assert len(slices) == 6 * len(panel) and not slices.duplicated(["organism", "candidate_antibiotic"]).any()
        for row in slices.itertuples(index=False):
            organism = test[test.organism == row.organism]
            group = organism[organism.candidate_antibiotic == row.candidate_antibiotic]
            denominator = len(organism[event].drop_duplicates())
            n, n_r = len(group), group.label.eq(0).sum()
            missed = (group.label.eq(0) & group[pcol].ge(threshold_s)).sum()
            for field, expected in {"test_organism_events": denominator, "observed_drug_events": n, "n": n,
                "n_r": n_r, "n_s": n - n_r, "missed_r": missed,
                "observed_panel_coverage": n / denominator if denominator else np.nan,
                "prevalence_r": n_r / n if n else np.nan,
                "recall_r": (n_r - missed) / n_r if n_r else np.nan}.items():
                close(getattr(row, field), expected, f"organism slice/{field}")
            assert row.supported_error_slice == bool(n >= 50 and n_r >= 20)
        record["sections"]["retained_label_coverage"] = {"training_selected_organisms": 6,
            "observed_or_absent_cells_checked": len(slices), "unobserved_outcomes_imputed": False}

        figure_index = json.loads((advanced / "figure_index.json").read_text(encoding="utf-8"))
        assert len(figure_index) == 8 and len({f["stem"] for f in figure_index}) == 8
        figure_paths = []
        for figure in figure_index:
            assert set(figure["files"]) == {"png"}
            for source in figure["source_tables"]:
                assert (advanced / source).is_file(), f"Figure source absent: {source}"
            for kind, name in figure["files"].items():
                path = advanced / name
                assert path.is_file() and path.stat().st_size > 1000, f"Figure export absent/empty: {name}"
                assert path.suffix.lower() == ".png"
                with Image.open(path) as image:
                    image.load()
                    assert image.format == "PNG", f"Wrong figure format: {name}"
                    assert min(image.size) >= 1000, f"Figure raster too small: {name}"
                figure_paths.append(name)
        for folder in (root / "results").rglob("figures"):
            if folder.is_dir():
                assert all(path.suffix.lower() == ".png" for path in folder.rglob("*") if path.is_file()), \
                    f"Non-PNG file in figure folder: {folder}"
        record["sections"]["figures"] = {"figures": len(figure_index), "valid_exports": len(figure_paths),
            "formats": ["PNG"]}

        report_skip = {"status": "skipped",
                       "reason": "Report checks skipped per repository scope; enable explicitly with --include-reports."}
        if include_reports:
            report = PdfReader(root / "report/report.pdf")
            report_text = [page.extract_text() or "" for page in report.pages]
            reference_pages = [i for i, text in enumerate(report_text) if "References and file guide" in text]
            assert reference_pages == [5] and len(report.pages) == 6
            assert "\u2014" not in "\n".join(report_text)
            report_meta = json.loads((root / "report/report_validation.json").read_text(encoding="utf-8"))
            assert report_meta["body_pages"] == 5 and report_meta["pdf_pages"] == 6
            record["sections"]["main_report"] = {"status": "passed", "body_pages": 5,
                                                   "reference_pages": 1, "em_dashes": 0}
        else:
            record["sections"]["main_report"] = report_skip.copy()

        optional = {"workbook": advanced / "advanced_results.xlsx"}
        optional_results = {}
        if include_reports:
            optional.update({"appendix_pdf": root / "report/advanced_analysis.pdf",
                             "gallery_html": root / "report/advanced_analysis.html"})
        else:
            optional_results.update({"appendix_pdf": report_skip.copy(), "gallery_html": report_skip.copy()})
        for name, path in optional.items():
            if not path.exists():
                assert not require_present, f"Required final artifact missing: {path.relative_to(root)}"
                optional_results[name] = {"status": "not yet present"}
                continue
            if name == "appendix_pdf":
                pdf = PdfReader(path)
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
                assert len(pdf.pages) == 10 and "\u2014" not in text
                optional_results[name] = {"status": "passed", "pages": len(pdf.pages), "em_dashes": 0}
            elif name == "gallery_html":
                text = path.read_text(encoding="utf-8")
                assert "\u2014" not in text
                for figure in figure_index:
                    assert figure["stem"] in text, f"Figure absent from gallery: {figure['stem']}"
                local_links = []
                for _, target in re.findall(r'(href|src)=[\"\']([^\"\']+)[\"\']', text):
                    if target.startswith(("#", "http:", "https:", "data:", "mailto:")):
                        continue
                    resolved = path.parent / target.split("#")[0]
                    assert resolved.exists(), f"Broken gallery relative link: {target}"
                    local_links.append(target)
                optional_results[name] = {"status": "passed", "indexed_figures": 8, "local_links": len(local_links), "em_dashes": 0}
            else:
                from openpyxl import load_workbook
                book = load_workbook(path, read_only=True, data_only=True)
                assert len(book.sheetnames) >= 10
                assert all(sheet.max_row > 0 for sheet in book.worksheets)
                index_rows = list(book["Index"].values)
                assert index_rows[0] == ("Sheet", "Rows", "Purpose", "Source file")
                csv_mapping = {row[3]: row for row in index_rows[1:] if row[0] in book.sheetnames and str(row[3]).endswith(".csv")}
                expected_csvs = {source.name for source in advanced.glob("*.csv") if "bootstrap_replicates" not in source.name}
                assert set(csv_mapping) == expected_csvs, "Workbook does not cover exactly the final non-bootstrap CSV tables."
                exported_rows = 0
                for source_name, mapping in csv_mapping.items():
                    sheet_name, n_rows, _, _ = mapping
                    source = read(source_name)
                    values = book[sheet_name].values
                    columns = next(values)
                    assert list(columns) == source.columns.tolist()
                    exported = pd.DataFrame(values, columns=columns)
                    assert len(exported) == len(source) == n_rows
                    pd.testing.assert_frame_equal(exported, source, check_dtype=False,
                                                  check_exact=False, rtol=1e-12, atol=1e-12)
                    exported_rows += len(exported)
                optional_results[name] = {"status": "passed", "sheets": book.sheetnames,
                    "csv_tables_matched_cell_for_cell": len(csv_mapping), "exported_data_rows": exported_rows}
                book.close()
        record["sections"]["supplementary_deliverables"] = optional_results
        record["independent_numeric_checks"] = len(comparisons)
        record["maximum_absolute_difference"] = max(comparisons, default=0.)
        record["status"] = "passed"
    except Exception as error:
        record["status"] = "failed"
        record["error"] = f"{type(error).__name__}: {error}"
        raise
    finally:
        (advanced / "validation.json").write_text(json.dumps(record, indent=2, allow_nan=False), encoding="utf-8")
    return record


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--require-present", action="store_true",
                        help="Require the workbook; also require appendix and gallery when --include-reports is set.")
    parser.add_argument("--include-reports", action="store_true",
                        help="Also check the optional main report, appendix and gallery in report/.")
    args = parser.parse_args()
    result = validate(args.root, args.require_present, args.include_reports)
    print(json.dumps({key: result[key] for key in ["status", "independent_numeric_checks", "maximum_absolute_difference"]}, indent=2))


if __name__ == "__main__":
    main()
