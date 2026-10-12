"""Build schema deliverables from the verified source and feature contracts."""
from pathlib import Path
import csv
import html
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SOURCE = "https://datadryad.org/dataset/doi:10.5061/dryad.jq2bvq8kp"
ARTICLE = "https://www.nature.com/articles/s41597-025-05649-7"
COHORT = "microbiology_cultures_cohort.csv"
DEMOGRAPHICS = "microbiology_cultures_demographics.csv"
WARD = "microbiology_cultures_ward_info.csv"
PENDING = "Pending: instructor extract and field mapping not provided"
VERIFIED_PATH = ROOT / "data" / "raw" / "verified_downloads.json"
VERIFIED = {x["file"]: x for x in json.loads(VERIFIED_PATH.read_text(encoding="utf-8"))} if VERIFIED_PATH.exists() else {}
SOURCE_STATUS = "Derived field; see source mapping and construction audit"
SPEC_PATH = ROOT / "src" / "armd" / "support_feature_spec.json"
SPEC = json.loads(SPEC_PATH.read_text(encoding="utf-8")) if SPEC_PATH.exists() else {"features": [], "excluded_tables": []}


def field(name, table, source_field, dtype, role, description, transform,
          missing, timing, allowed="", required="yes"):
    return {
        "canonical_field": name,
        "armd_source_table": table,
        "armd_source_field": source_field,
        "logical_dtype": dtype,
        "role": role,
        "description": description,
        "transform_or_rule": transform,
        "missing_handling": missing,
        "availability_at_prediction": timing,
        "allowed_values_or_units": allowed,
        "present_in_analytical_contract": required,
        "instructor_source_mapping": PENDING,
        "verification_status": ("Local bytes match Dryad version 403223 SHA-256" if table in VERIFIED else SOURCE_STATUS),
    }


def dictionary():
    rows = [
        field("patient_id", COHORT, "anon_id", "string", "linkage and split only",
              "Anonymous patient identifier; excluded from predictors.", "Trim; retain as string.",
              "Exclude rows without a patient key.", "Audit use only", "Nonempty anonymous code"),
        field("culture_id", COHORT, "order_proc_id_coded", "string", "linkage only",
              "Culture order identifier; use with patient_id.", "Trim; retain as string.",
              "Exclude rows without an event key.", "Audit use only", "Nonempty anonymous code"),
        field("organism", COHORT, "organism", "string", "grain and subgroup only",
              "Current organism for the requested row grain; excluded from predictors at culture order.",
              "Trim, lowercase, collapse repeated spaces; do not invent species mappings.",
              "Exclude rows without organism.", "Usually identified after order; never a primary predictor"),
        field("candidate_antibiotic", COHORT, "antibiotic", "string category", "predictor and grain",
              "Candidate drug whose susceptibility is predicted.",
              "Trim and lowercase; standardize named aliases; restrict to configured candidate list.",
              "Exclude absent or out-of-list candidate names.", "Candidate can be selected at order time",
              "One configured candidate name"),
        field("culture_time", COHORT, "order_time_jittered_utc", "UTC timestamp", "audit only",
              "Patient-shifted culture order time. It is not a real calendar date.",
              "Parse as UTC. Exclude events with conflicting nonmissing time strings or culture types.",
              "Exclude missing or unparseable order times.", "Prediction reference time; not a predictor", "Shifted UTC date and time"),
        field("culture_type", COHORT, "culture_description", "string category", "predictor",
              "Culture specimen or order description.",
              "Trim text; retain source category. Do not infer a cross-source mapping.",
              "Keep missing; encoder uses a missing category. Exclude conflicting event metadata.",
              "Order information available at culture order"),
        field("label", COHORT, "susceptibility", "integer", "outcome only",
              "Susceptible is 1; resistant is 0.",
              "Accept S/Susceptible and R/Resistant ignoring case. Drop the entire key if any label is invalid or if both 0 and 1 occur. Collapse agreeing repeats.",
              "Exclude intermediate, unknown, inconclusive, synergy, missing and all other labels.",
              "Future laboratory result; never a predictor", "0 or 1"),
        field("age_group", DEMOGRAPHICS, "age", "string category", "predictor",
              "Source age band. Do not reconstruct exact age or assume a meaning for the oldest band.",
              "Trim; replace long dash characters with hyphens; exclude bands with an explicit lower bound below 18.",
              "Unmatched, absent or conflicting values stay missing; encode a missing category.",
              "Demographics documented at order time", "Source age bands"),
        field("gender_code", DEMOGRAPHICS, "gender", "string category", "predictor",
              "Anonymous gender code. The source does not disclose which code means which gender.",
              "Trim; retain source code. Do not label codes as male or female.",
              "Unmatched, absent or conflicting values stay missing; encode a missing category.",
              "Demographics documented at order time", "Source code, typically 0 or 1"),
    ]
    for code, meaning in [("ip", "inpatient"), ("op", "outpatient"),
                          ("er", "emergency department"), ("icu", "intensive care")]:
        rows.append(field(
            "ward_" + code, WARD, "hosp_ward_" + code.upper(), "nullable binary", "predictor",
            "Flag for the " + meaning + " setting. Ward flags can overlap and are not hospital identifiers.",
            "Collapse agreeing repeats; conflicting values become missing. Map true/false to 1/0; other values become missing.",
            "Absent is missing, never zero; encode an explicit __MISSING__ category.",
            "Setting at culture collection per source; assumed available at order, not individually time-verified", "0, 1 or missing"))
    rows.extend([
        field("dataset", "config.json", "source_dataset", "string category", "provenance and subgroup only",
              "Source dataset name.", "Copy the configured source label.", "Required configuration value.",
              "Audit use only", "ARMD for this mapping"),
        field("split", "derived", "patient_id and config.seed", "string category", "partition only",
              "Patient-level allocation shared by all cultures, organisms and drugs for that patient.",
              "SHA-256 of split seed and patient key; fractions below 0.6 train, below 0.8 validation, otherwise test.",
              "No missing value is allowed.", "Set before fitting any preprocessing or model",
              "train, validation or test"),
        field("site_id", "no verified ARMD source", "not available in published ARMD schema", "nullable string", "optional subgroup only",
              "A real hospital identifier, if an approved future extract supplies one. Do not substitute ward for site.",
              "Preserve a supplied site field for audit; no site is derived for ARMD.",
              "ARMD site is unavailable. Do not create pseudo-hospital labels.", "Not a predictor",
              "Verified hospital code only", required="optional; absent in documented ARMD input"),
    ])
    for item in SPEC["features"]:
        rows.append(field(item["name"], item["source_table"], "; ".join(item["source_columns"]),
                          item["dtype"], item["model_role"], item["description"],
                          item["aggregation"], item["missingness"], item["timing"], item["units"]))
    return rows


def harmonization(data_dictionary):
    rows = [dict(x) for x in data_dictionary]
    excluded = [
        ("encounter_id", COHORT, "pat_enc_csn_id_coded", "Identifier is not a predictor; culture linkage uses patient and culture order."),
        ("culture_positive_flag", COHORT, "was_positive", "Used only to exclude explicit negative cultures; not a predictor and not kept in final rows."),
        ("current_susceptibility", COHORT, "susceptibility", "Outcome source only. Current panels, MICs and drug-result combinations cannot be predictors."),
        ("prior_resistance", "microbiology_cultures_microbial_resistance.csv", "organism; antibiotic; resistant_time_to_culturetime", "Requested for a possible extension, excluded from current core. Confirm the time sign, strict pre-order result availability and event aggregation before use."),
        ("prior_organism", "microbiology_culture_prior_infecting_organism.csv", "prior_organism; prior_infecting_organism_days_to_culture", "Excluded. Needs verified prior result timing and a defined history window."),
        ("prior_medication", "microbiology_cultures_prior_med.csv", "medication_name; medication_time_to_culturetime; medication_category", "Excluded. Requires a time sign check, a pre-order cutoff and one event summary."),
        ("antibiotic_class_exposure", "microbiology_cultures_antibiotic_class_exposure.csv", "antibiotic_class; time_to_culturetime", "Requested for a possible extension, excluded from current core. Common-source availability is unknown and prior exposure timing is not checked."),
        ("antibiotic_subtype_exposure", "microbiology_cultures_antibiotic_subtype_exposure.csv", "antibiotic_subtype; time_to_culturetime", "Excluded. Common-source availability is unknown and prior exposure timing is not checked."),
        ("laboratory_history", "microbiology_cultures_labs.csv", "Period_Day; first_wbc; last_wbc; first_cr; last_cr (examples)", "Requested for a possible extension, excluded from current core. Source describes a pre-order window, but units, exact summary columns and per-extract timing need validation."),
        ("vital_history", "microbiology_cultures_vitals.csv", "first_heartrate; last_heartrate; first_temp; last_temp (examples)", "Requested for a possible extension, excluded from current core. Units, source windows and exact summary schema need validation."),
        ("comorbidity_history", "microbiology_cultures_comorbidity.csv", "comorbidity_component; comorbidity_component_start_days_culture; comorbidity_component_end_days_culture", "Excluded from scaled core. A missing end date can mean ongoing disease; it cannot mean no disease."),
        ("prior_procedure", "microbiology_cultures_priorprocedures.csv", "procedure_name; procedure_time_to_culturetime", "Excluded. Needs strict pre-order timing and one summary per event."),
        ("area_deprivation", "microbiology_cultures_adi_scores.csv", "adi_score; adi_state_rank", "Excluded. Common-source coverage, geographic meaning and timing are not verified."),
        ("nursing_home_history", "microbiology_cultures_nursing_home_visits.csv", "nursing_home_visit_culture", "Excluded. Missing linked records do not establish no nursing home history."),
        ("implied_label", "microbiology_cultures_implied_susceptibility.csv", "Implied_susceptibility", "Excluded from both primary labels and features; inferred labels change the target and can create circular evaluation."),
        ("implied_label_rule", "implied_susceptibility_rules.csv", "Organism; Antibiotic; Rule", "Excluded from primary construction. Any future inferred-label analysis needs separate reporting."),
    ]
    used_support = {x["source_table"] for x in SPEC["features"]}
    exclusions_by_table = {x["source_table"]: x["reason"] for x in SPEC["excluded_tables"]}
    for name, table, col, reason in excluded:
        if table in used_support:
            continue
        reason = exclusions_by_table.get(table, reason)
        rows.append(field(name, table, col, "source dependent", "excluded or construction-only",
                          reason, "Not exported as a predictor.", "No absence-to-zero conversion.",
                          "No approved feature timing contract", required="no"))
    return rows


def source_files():
    choices = {
        COHORT: ("include", "Primary event, organism, candidate and observed result table."),
        DEMOGRAPHICS: ("include when available", "Age band and anonymous gender code; reduce to one patient-event row before joining."),
        WARD: ("include when available", "Four care-setting flags; reduce to one patient-event row before joining."),
        "README.md": ("documentation only", "Publisher schema and source definitions. This is not a data table."),
        "implied_susceptibility_rules.csv": ("exclude", "The primary task uses observed labels. Rules are not model inputs."),
        "microbiology_cultures_implied_susceptibility.csv": ("exclude", "Inferred labels are separate from measured primary outcomes."),
        "microbiology_cultures_microbial_resistance.csv": ("pending extension; not used", "Requested file. Historical resistance requires a strict cutoff and timing validation."),
        "microbiology_culture_prior_infecting_organism.csv": ("exclude", "Prior organism history needs a verified result time and aggregation."),
        "microbiology_cultures_prior_med.csv": ("exclude", "Medication timing and common-source availability are not verified."),
        "microbiology_cultures_antibiotic_class_exposure.csv": ("pending extension; not used", "Requested file. Exposure timing and common-source availability are not verified."),
        "microbiology_cultures_antibiotic_subtype_exposure.csv": ("exclude", "Exposure timing and common-source availability are not verified."),
        "microbiology_cultures_labs.csv": ("pending extension; not used", "Requested file. Needs unit, window and schema checks before use."),
        "microbiology_cultures_vitals.csv": ("pending extension; not used", "Requested file. Needs unit, window and schema checks before use."),
        "microbiology_cultures_comorbidity.csv": ("exclude", "Large historical table outside scaled core; requires onset and ongoing-condition rules."),
        "microbiology_cultures_priorprocedures.csv": ("exclude", "Historical timing and event summaries are not yet validated."),
        "microbiology_cultures_adi_scores.csv": ("exclude", "Common-source availability and geographic comparability are unverified."),
        "microbiology_cultures_nursing_home_visits.csv": ("exclude", "Coverage and prior-visit timing require validation."),
    }
    manifest_path = ROOT / "data" / "raw" / "dryad_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else {}
    files = {x["path"]: x for x in manifest.get("_embedded", {}).get("stash:files", [])}
    for item in SPEC["features"]:
        choices[item["source_table"]] = ("include selected summaries", item["timing"] + "; " + item["missingness"])
    for item in SPEC["excluded_tables"]:
        choices[item["source_table"]] = ("exclude from model", item["reason"])
    rows = []
    for name, (decision, reason) in choices.items():
        item = files.get(name, {})
        rows.append({
            "file": name, "kind": "documentation" if name == "README.md" else "CSV data table",
            "decision": decision, "reason": reason,
            "manifest_version": "403223 (2025-10-22)",
            "publisher_size_bytes": item.get("size", "unavailable"),
            "publisher_sha256": item.get("digest", "unavailable"),
            "source_url": SOURCE,
            "verification_status": ("Local bytes match Dryad version 403223 SHA-256" if name in VERIFIED else "Not supplied; not used"),
        })
    return rows


def write_csv(path, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def write_workbook(path, sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        sheet.append(list(rows[0]))
        for row in rows:
            sheet.append(list(row.values()))
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="173E57")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 36 if col > 1 else 29
        sheet.row_dimensions[1].height = 42
        for row in range(2, sheet.max_row + 1):
            sheet.row_dimensions[row].height = 86
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
    workbook.save(path)


def table(rows, keys):
    return "<table><thead><tr>" + "".join("<th>" + html.escape(k.replace("_", " ")) + "</th>" for k in keys) + "</tr></thead><tbody>" + "".join(
        "<tr>" + "".join("<td>" + html.escape(str(row[k])) + "</td>" for k in keys) + "</tr>" for row in rows) + "</tbody></table>"


def methods(config, inventory):
    drugs = ", ".join(config["antibiotics"])
    numeric = ", ".join(x["name"] for x in SPEC["features"])
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>ARMD construction and evaluation protocol</title>
<style>body{{font:16px/1.55 Arial,sans-serif;color:#202c36;max-width:1000px;margin:40px auto;padding:0 24px}}h1,h2{{line-height:1.2;color:#173e57}}h1{{font-size:30px}}h2{{font-size:23px;margin-top:30px}}table{{border-collapse:collapse;width:100%;font-size:13px}}th,td{{border:1px solid #bbc8d1;text-align:left;vertical-align:top;padding:8px}}th{{background:#e8f0f4}}code{{background:#f0f3f5;padding:2px 4px}}.status{{border-left:5px solid #c88915;background:#fff7e7;padding:14px 20px}}a{{color:#145377}}@media print{{body{{font-size:10pt;margin:0;max-width:none}}h2{{break-after:avoid}}tr{{break-inside:avoid}}.status{{background:white}}}}</style>
</head><body>
<h1>ARMD construction and evaluation protocol</h1>
<p class="status"><strong>Status:</strong> These notes explain the methods and data fields. The supplied ARMD source files match the pinned publisher checksums. The real-data results are saved in the results folder. The instructor extract and exact candidate list were not supplied. The current run uses the stated provisional ARMD panel. No synthetic substitute is used for the main analysis.</p>

<h2>1. The prediction task</h2>
<p>At the time a culture is ordered, estimate whether a candidate antibiotic will have a susceptible result. Each analytical row is one patient, culture order, organism and candidate antibiotic. The label is 1 for susceptible and 0 for resistant. This is a prediction study, not a treatment recommendation.</p>
<p>Only records with an observed S or R label enter the primary target. An untested candidate has no known label. The pipeline does not create a full nine-drug panel for every culture and does not turn an untested drug into resistant. Organism is needed to define a row and examine subgroups. It is not a predictor because the current organism is usually not known when the culture is ordered.</p>
<p>Eligible rows need a patient key, culture key, organism, parseable shifted order time and a configured candidate. Explicit negative-culture rows are removed. ARMD is described as an adult cohort. As an extra check, an age band with an explicit lower bound below 18 is removed. Missing age is retained and counted; the pipeline does not infer that a person with no age record is an adult. This relies on the source's adult inclusion rule.</p>
<p>The current nine-drug list is <strong>provisional</strong>: {html.escape(drugs)}. This is a configurable teaching list. It has not been verified as the standardized STEWARD-AW list. The analysis can be rerun if the instructor supplies a different intended list.</p>

<h2>2. Resolve repeated and unclear records</h2>
<p>The unique key is <code>patient_id, culture_id, organism, candidate_antibiotic</code>. Text is trimmed and declared null strings become missing. Named drug aliases are mapped to one name. Organism text is lowercased and repeated spaces are removed, without guessing a species.</p>
<p>For each key, matching S labels or matching R labels collapse to one row. If both S and R occur, the whole key is dropped. If any row has an intermediate, unknown, blank, inconclusive, synergy or other invalid label, the whole key is dropped, even when another row says S or R. This conservative rule also drops S plus I. Counts and reasons are written to the label audit.</p>
<p>Order time and culture type must agree within an event. Events with conflicting nonmissing values are removed before candidate filtering. This prevents an arbitrary first row from setting the event metadata. Time values are checked before parsing, so alternate string formats for the same instant can also cause exclusion; the audit makes that loss visible.</p>

<h2>3. Select a small common feature set</h2>
<p>The eight categorical predictors are candidate antibiotic, age band, anonymous gender code, culture type and four ward flags. The {len(SPEC['features'])} numeric predictors are {html.escape(numeric)}. These fields use the published ARMD schema and the checked source files. The instructor-source mapping is still pending. They are a proposed shared set, not a verified intersection of several datasets.</p>
<p>Laboratory summaries cover the 14 days before culture order, and vital signs cover the prior 48 hours according to the author manuscript. Exact release times are not supplied. Laboratory and vital units are not fully documented, so values keep their source units and no clinical unit conversion is guessed. Only selected median summaries are retained. Repeated event summaries must agree; conflicting medians become missing instead of being averaged. The feature dictionary gives the exact timing windows and aggregation for each history feature.</p>
<p>Age bands are categories, not exact ages. The meanings of gender codes 0 and 1 are not disclosed, so they stay anonymous. Ward flags can overlap. They do not identify the two hospitals. A missing ward record means unknown setting, not absence from all wards.</p>
<p>Demographics are documented at culture order. Ward flags describe where the culture was collected. The primary design assumes that setting is available at order time. A separate ward recording timestamp is not in the documented core, so this assumption cannot be checked for each row. A future extract should verify it or omit these flags.</p>

<h2>4. Prepare supporting tables before joining</h2>
<p>Demographics and ward records are first reduced to one row per patient and culture order. Matching nonmissing values collapse. If one feature has several different nonmissing values, only that feature becomes missing. No majority vote is used. Each support table is then left joined to the cohort using both keys, with an explicit many-to-one check.</p>
<p>A left join keeps an eligible culture when a support record is absent. The pipeline checks that joins do not change the number of cohort rows and that the final analytical key is unique. Raw identifiers remain in the analytical file for these checks and patient grouping. The model receives only the named feature list.</p>

<h2>5. Split first, then learn from training data</h2>
<p>A stable SHA-256 hash samples whole patients using seed {config['seed']}. The current sampling fraction is {config['sample_fraction']:.0%}. A separate hash gives each patient about 60% chance of training, 20% validation and 20% testing. The realized row and patient fractions can differ. Every culture, organism and candidate for a patient stays in that patient's split.</p>
<p>Fixed label rules, field renaming and key-local agreement checks do not learn population statistics. After those construction steps, the split is assigned before fitting encoders, imputers, scaling, class weights or models. Any categories, averages, medians or class weights learned from data must use training data only. Validation and test rows keep their original class prevalence.</p>
<p>Current susceptibility, current MIC, current resistance summaries, current organism, identifiers and order dates are excluded from model inputs. Current drug panels can reveal outcomes and must not be added as features. Historical drug exposure and prior candidate resistance use a strictly earlier cutoff and the documented day-difference direction. Same-day and future history are excluded. Missing linked history stays missing. Prior candidate resistance is matched to the candidate, so its intended join key also includes that drug.</p>
<p>ARMD applies patient-specific date shifts. These preserve timing within a patient but do not establish real calendar order across patients. A split using shifted dates is not a real prospective time test. Patient splitting limits patient overlap; it does not prove future or outside-hospital performance.</p>

<h2>6. Model and evaluation contract</h2>
<p>The models are logistic regression and LightGBM, each fitted both with and without balanced class weights. There are also two simple baselines: the overall training susceptible share and the susceptible share for each candidate. An unseen candidate falls back to the overall share. The four learned variants have fixed settings rather than a large hyperparameter search.</p>
<p>The training patients are split again: 80% for base-model fitting and 20% for calibration. This division is patient-disjoint and must contain both labels in each part. Missing categorical values get the explicit category <code>__MISSING__</code>. A one-hot encoder learns categories from base-model fitting rows only; it ignores unseen categories when scoring. Ward zero stays separate from a missing ward value. Numeric fields use medians learned from the fitting subset, missingness flags and scaling fitted only on that subset. Numeric fields entirely absent in fitting receive a fixed placeholder plus an explicit missing indicator; this is not an observed measurement of zero. The source and analytical files keep missing values.</p>
<p>Logistic regression uses L2 regularization with <code>C=1</code>, the lbfgs solver and at most 1,000 iterations. LightGBM uses 220 trees, learning rate 0.04, 15 leaves, at least 50 rows per leaf, L2 regularization 1 and column fraction 0.9. Balanced weights are <code>n / (2 * class_count)</code>, computed only on base-model fitting rows. Both baselines use that same fitting subset. Validation, calibration and test rows keep their original class prevalence.</p>
<p>Each learned model has a sigmoid calibration model fitted on the separate 20% training calibration subset. It is an unweighted logistic regression of the label on the logit of the base probability, with <code>C=1000000</code>. Weighted fitting can distort probabilities, so calibrated and raw probability diagnostics are both reported. Validation is not used to fit calibration.</p>
<p>Each model's threshold maximizes balanced accuracy on validation data; ties use the threshold nearest 0.5. The selected learned model has the highest validation antibiotic-macro resistant average precision, with lower Brier score as the tie-break. These choices are frozen before test scoring. Test results for all pre-specified models remain visible, but test performance does not select the winner. Diagnostic calibration slope and intercept fitted on test predictions describe error only; they do not change test predictions.</p>
<p>Let <code>p</code> be the predicted probability of susceptibility and <code>y</code> be the observed label. Resistance scores are <code>1-p</code>. At a susceptibility threshold <code>t</code>, predict susceptible when <code>p &gt;= t</code> and resistant otherwise. A resistant false negative means a resistant row was predicted susceptible. Report the chosen threshold and how it was set.</p>
<ul>
<li><strong>AUROC:</strong> the chance a positive row receives a higher score than a negative row. Report susceptible and resistant views with the correct score direction.</li>
<li><strong>AUPRC:</strong> average precision. State this numerical definition because trapezoidal precision-recall area differs. Report both classes. Resistant-class average precision uses resistant rows as positive and <code>1-p</code> as the score.</li>
<li><strong>Class-macro metrics:</strong> an unweighted average over the susceptible and resistant class. Binary class-macro AUROC equals either properly oriented class AUROC.</li>
<li><strong>Drug-macro metrics:</strong> an unweighted average of each eligible candidate's metric. Name these separately from class-macro metrics. Report how many drugs were eligible. A candidate with only one observed class has no defined AUROC and must not silently receive zero.</li>
<li><strong>Balanced accuracy:</strong> the mean of susceptible recall and resistant recall.</li>
<li><strong>Resistant precision:</strong> the share of predicted resistant rows that are resistant. <strong>Resistant recall:</strong> the share of resistant rows detected. Precision is set to zero when no row is predicted resistant. Recall is missing when there are no resistant rows.</li>
<li><strong>Brier score:</strong> the mean of <code>(p-y)^2</code>; lower is better. It tests probability error and depends on class prevalence.</li>
<li><strong>Calibration plot:</strong> compare mean predicted susceptibility with the observed susceptible share in probability bins. Show a diagonal reference, bin counts and the probability distribution when practical. Sparse bins give uncertain estimates.</li>
</ul>
<p>Report overall, per-candidate, culture-type, age-band, gender-code, source and shifted five-year-era results, with sample sizes and class counts. Groups below 100 rows or with fewer than 20 rows in either class are flagged as small. Single-class groups have missing AUROC and average precision under this reporting convention. Resistant is not always the minority. Shifted calendar eras are descriptive, not evidence of real calendar change.</p>
<p>Add hospital results only if a real hospital identifier exists. With documented ARMD input, there is one source label and no verified hospital field, so independent dataset or site validation is unavailable. Ward analysis cannot meet that requirement.</p>
<p>Patient-cluster bootstrap intervals respect repeated rows within a patient. Patients are sampled with replacement and all their rows are weighted by their draw count. All models use the same draws, so paired differences use the same sampled patients. The configuration requests {config['bootstrap_replicates']} replicates. Percentile intervals from a small bootstrap are approximate. Undefined replicate metrics are skipped and usable counts are reported. Models are not refitted in the bootstrap, so these intervals describe test-sample uncertainty for the fitted models, not uncertainty from retraining. Patient clustering does not account for every shared hospital or calendar effect.</p>

<h2>7. Profiles and files to inspect</h2>
<p>The cohort summary reports rows, patients, culture events, organism-events, shifted time span, label prevalence, candidate depth and duplicate keys, overall and by split. Candidate depth is the number of retained candidate antibiotics for one patient, culture and organism. It counts observed eligible labels, not all drugs that could have been selected.</p>
<p>Missingness is reported for every field overall and by split. Class balance is reported by split and candidate. Construction logs show dropped labels, repeated keys, support conflicts and missing tables. The model package adds model comparison, group results, calibration and uncertainty outputs. Files can be empty or absent when a run is blocked; that is not evidence of perfect data or model performance.</p>
<p>The final schema contract is in <a href="final_data_dictionary.csv">final_data_dictionary.csv</a>. The <a href="feature_harmonization_matrix.xlsx">Excel workbook</a> includes the harmonization matrix, dictionary, source-file decisions and target list. The CSV matrix names fields excluded from the core as well as retained fields. Source mappings are documentation, not proof that every source value matches the contract.</p>

<h2>8. All published source files</h2>
<p>The pinned Dryad manifest has 17 files: 16 CSV data tables and one README. The model uses selected fields from the cohort, demographics, ward, laboratory, vital-sign, antibiotic-exposure and resistance-history tables. The inventory states why each remaining table is excluded. The supplied files are checked against the publisher's SHA-256 digests.</p>
{table(inventory, ['file', 'decision', 'reason'])}

<h2>9. What the results would and would not show</h2>
<p>Only tested drugs with usable S or R results are labeled. Clinicians and laboratories decide what gets tested. This can make the labeled sample different from all patients, all organisms or all possible drug choices. Dropping intermediate and conflicting labels can change prevalence too. A strong test score on observed labels does not establish the same score for untested candidates.</p>
<p>Rows with the same patient, event and organism can occur for several candidates. A random row split would leak shared patient information across sets. The patient split prevents that overlap, while cluster intervals reflect some of the repeated-row dependence. Multi-organism cultures can still produce hard cases because the pre-order inputs cannot distinguish organisms within one event.</p>
<p>The source spans changing laboratory practices and antibiotic use. This scaled model uses a small feature set. It cannot validate causal treatment effects, patient outcomes, present-day breakpoints or performance in another hospital. No comparison between real hospitals is possible without a hospital field. A synthetic run can test code paths but cannot estimate clinical performance or stand in for an approved real-data study.</p>

<h2>10. Source and writing notes</h2>
<p>Schema names and file decisions use the <a href="{SOURCE}">ARMD Dryad record and README</a>. The study context comes from the <a href="{ARTICLE}">published ARMD article</a> and its <a href="https://arxiv.org/pdf/2503.07664">author manuscript</a>. The pinned manifest identifies version 403223, dated 22 October 2025. The download script checks publisher file sizes and SHA-256 digests before accepting downloaded source files.</p>
<p>These notes follow the <a href="https://www.archives.gov/open/plain-writing/10-principles.html">National Archives plain-language principles</a>: short sentences, everyday words and a clear main point. The review of <a href="https://libguides.usc.edu/writingguide/AIwriting">USC Libraries guidance on AI writing</a> informed checks for false facts, invented citations and unsupported claims. Numeric results come from saved outputs. AI assistance was used for code and writing.</p>

<h2>11. Exploratory follow-up analysis</h2>
<p>The <a href="../results/advanced/figures/">extra figures</a> extend the original study. The <a href="../results/advanced/figure_index.json">figure notes</a> explain their scope and source tables, and the <a href="../results/advanced/advanced_results.xlsx">results workbook</a> contains the supporting numbers and method details. These checks were added after the first test results were known. The original fitted model and results stay fixed. A saved protocol fingerprints seven original artifacts. New findings need independent confirmation.</p>
<p>Feature-group comparisons reuse the original model-fitting and calibration patients, model settings and preprocessing rules. Only validation outcomes score these comparisons. The groups are basic fields, basic fields plus laboratory and vital summaries, basic fields plus prior history, all features, and basic fields plus information-availability patterns. The last control replaces observed numeric values with a fixed constant while keeping missing indicators. It does not replace missing clinical values with zero. A known prior-resistance record is itself informative, so this control does not isolate recording practice alone.</p>
<p>Drug-level checks use frozen test predictions and the original cutoff. Probability figures use resistance probability, which equals one minus predicted susceptibility. Overall reliability bins use ten validation quantiles; drug-level bins use five. Duplicate edges collapse, and all bin boundaries are fixed before test evaluation. Bins below 100 rows or 20 resistant results are marked sparse. Brier error, bin agreement and probability distributions are shown together.</p>
<p>Three further cutoffs aim for 70%, 80% and 90% pooled resistant recall on validation. Each keeps ties together and selects the highest resistance cutoff that reaches the target. The exact susceptible-probability rule is stored in JSON, avoiding rounding changes at tied scores. These cutoffs then apply unchanged to test. False alerts and missed resistance are expressed per 1,000 candidate rows, not per patient. A pooled target does not guarantee a target for every drug and is not a clinical treatment rule.</p>
<p>New intervals use 500 shared patient-bootstrap samples. All rows for a sampled patient stay together. Feature-group differences and before/after calibration comparisons use paired draws. Models, cutoffs and bin edges remain fixed during resampling. Intervals are pointwise and omit model-training and selection uncertainty. An interval fixed at zero or one only describes the observed sample. The model-comparison figure keeps the original 200-draw intervals.</p>
<p>The coverage display selects six organisms by their training event frequency. Its denominator is test organism-events with at least one retained S/R result in the panel. Coverage measures retained label availability, not the raw testing rate. An absent label never becomes an outcome. New missingness plots count shared fields once per culture and candidate resistance history once per culture and candidate. The original row-level missingness table remains available.</p>
<p>Run <code>scripts/run_advanced_analysis.py</code> after the primary pipeline. The command builds the extra tables, eight figures, workbook and submission archive. The <code>--presentation-only</code> option reuses saved analysis results. Figure explanations stay in the figure notes and workbook. The independent validator recalculates estimates and checks protected hashes before packaging.</p>
<p>Reports are optional and are ignored by Git. Add <code>--include-reports</code> to <code>scripts/run_advanced_analysis.py --presentation-only</code> to generate them locally and include them in the archive. The separate <code>scripts/package_submission.py</code> command includes reports only when passed <code>--include-reports</code>.</p>
</body></html>"""


def main():
    DOCS.mkdir(parents=True, exist_ok=True)
    config = json.loads((ROOT / "config.json").read_text(encoding="utf-8"))
    schema_dictionary = dictionary()
    data_dictionary = [row for row in schema_dictionary if row["present_in_analytical_contract"] == "yes"]
    matrix = harmonization(schema_dictionary)
    inventory = source_files()
    targets = [{"candidate_antibiotic": x, "selection_status": config["target_list_status"],
                "source": "config.json", "verified_steward_aw_standard": "no"}
               for x in config["antibiotics"]]
    write_csv(DOCS / "feature_harmonization_matrix.csv", matrix)
    write_csv(DOCS / "final_data_dictionary.csv", data_dictionary)
    write_csv(DOCS / "source_file_decisions.csv", inventory)
    write_csv(DOCS / "target_antibiotics.csv", targets)
    sheets = {"Harmonization": matrix, "Data_dictionary": data_dictionary,
              "Source_files": inventory, "Target_antibiotics": targets}
    write_workbook(DOCS / "feature_harmonization_matrix.xlsx", sheets)
    write_workbook(DOCS / "final_data_dictionary.xlsx", {"Data_dictionary": data_dictionary})
    protocol = methods(config, inventory)
    if "\u2014" in protocol:
        raise ValueError("The protocol must not contain em dashes")
    (DOCS / "methods.html").write_text(protocol, encoding="utf-8")
    print(f"Wrote {len(matrix)} harmonization rows, {len(data_dictionary)} dictionary rows, "
          f"{len(inventory)} source-file decisions, and the method appendix.")


if __name__ == "__main__":
    main()
