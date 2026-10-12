"""Make compact exploratory figures and a reviewer-friendly results workbook."""
from pathlib import Path
import os
import json
import argparse
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
os.environ.setdefault('MPLCONFIGDIR', str(ROOT / '.cache' / 'matplotlib'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import pandas as pd
from armd.visual_style import clean_figure
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--results', type=Path, default=ROOT / 'results')
    args = parser.parse_args()
    out = args.results
    figures = out / 'figures'
    figures.mkdir(exist_ok=True)
    missing = pd.read_csv(out / 'missingness.csv')
    miss = missing[(missing['split']=='all') & missing['is_predictor']].sort_values('missing_fraction')
    balance = pd.read_csv(out / 'class_balance.csv')
    pooled = balance.groupby('candidate_antibiotic')[['n','resistant_n']].sum()
    pooled['resistant_fraction'] = pooled.resistant_n / pooled.n
    pooled = pooled.sort_values('resistant_fraction')
    plt.rcParams.update({'font.size': 9, 'axes.spines.top': False, 'axes.spines.right': False})
    fig, axes = plt.subplots(1, 2, figsize=(12, 5.5))
    axes[0].barh(miss.feature.str.replace('_',' '), miss.missing_fraction * 100, color='#287d8e')
    axes[0].set(xlabel='Missing (%)', ylabel='Input feature', xlim=(0,100))
    axes[1].barh(pooled.index, pooled.resistant_fraction * 100, color='#ba583c')
    axes[1].set(xlabel='Resistant (%)', ylabel='Candidate antibiotic', xlim=(0,100))
    clean_figure(fig)
    fig.tight_layout()
    fig.savefig(figures / 'missingness_and_resistance.png', dpi=180)
    plt.close(fig)
    depth = pd.read_csv(out / 'candidate_depth.csv')
    fig, ax = plt.subplots(figsize=(7,4))
    ax.bar(depth.candidate_depth, depth.organism_events, color='#287d8e')
    ax.set(xlabel='Observed candidate antibiotics per organism event', ylabel='Organism events',
           xticks=list(range(1,10)))
    clean_figure(fig)
    fig.tight_layout()
    fig.savefig(figures / 'candidate_depth.png', dpi=180)
    plt.close(fig)
    selected = json.loads((out / 'models/evaluation_summary.json').read_text())['selected_model']
    subgroup = pd.read_csv(out / 'models/subgroup_metrics.csv')
    perdrug = subgroup[(subgroup.model==selected) & (subgroup.group_field=='candidate_antibiotic')].sort_values('auprc_r')
    fig, ax = plt.subplots(figsize=(8,4.8))
    ax.barh(perdrug.group_value, perdrug.auprc_r, color='#287d8e')
    ax.scatter(1-perdrug.prevalence_s, perdrug.group_value, color='#ba583c', marker='|',s=160)
    ax.set(xlabel='Resistance average precision / resistance prevalence', ylabel='Candidate antibiotic', xlim=(0,1))
    clean_figure(fig)
    fig.tight_layout()
    fig.savefig(figures / 'per_antibiotic_performance.png', dpi=180)
    plt.close(fig)
    sheets = {
        'Cohort': out/'cohort_summary.csv', 'Missingness':out/'missingness.csv',
        'Class_balance':out/'class_balance.csv', 'Candidate_depth':out/'candidate_depth.csv',
        'Model_comparison':out/'models/model_comparison.csv',
        'Calibration':out/'models/calibration_diagnostics.csv',
        'Test_intervals':out/'models/test_confidence_intervals.csv',
        'Paired_differences':out/'models/paired_model_differences.csv',
        'Subgroups':out/'models/subgroup_metrics.csv',
        'Validation_importance':out/'models/validation_permutation_importance.csv',
    }
    with pd.ExcelWriter(out/'results_package.xlsx',engine='openpyxl') as writer:
        for name, file in sheets.items():
            if file.exists():
                pd.read_csv(file).to_excel(writer, sheet_name=name, index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes='A2'
            ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:
                cell.font=Font(color='FFFFFF',bold=True)
                cell.fill=PatternFill('solid',fgColor='17485B')
                cell.alignment=Alignment(wrap_text=True,vertical='top')
            ws.row_dimensions[1].height=38
            for column in ws.columns:
                width=max(len(str(c.value or '')) for c in list(column)[:100])
                ws.column_dimensions[get_column_letter(column[0].column)].width=min(max(width+2,14),40)
    print('Created results workbook and figures.', flush=True)


if __name__ == '__main__':
    main()
